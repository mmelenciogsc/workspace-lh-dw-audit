"""Dependency-aware Fabric item audit and guarded cleanup helpers."""

from __future__ import annotations

import argparse
import json
import os
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

import openpyxl


NO_DEPENDENCY_TEXT = re.compile(r"^no dependencies found$", re.I)
DEPENDENCY_COLUMNS = (
    "Direct semantic models",
    "Reports (via direct models)",
    "Shortcut into lakehouses (incl. chained shortcuts)",
    "Indirect semantic models (via shortcut)",
    "Indirect reports (via shortcut)",
    "Dataflows reading it",
    "Notebooks/pipelines/copy jobs reading it",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header_map(values: Iterable[Any]) -> dict[str, int]:
    return {_text(value).casefold(): index for index, value in enumerate(values, start=1) if _text(value)}


def _cell(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    position = headers.get(name.casefold())
    return row[position - 1] if position and position <= len(row) else None


def _item_parts(item: str) -> tuple[str, str]:
    name, separator, kind = item.rpartition(".")
    return (name, kind) if separator and kind.casefold() in {"lakehouse", "warehouse"} else (item, "")


@dataclass
class AuditDecision:
    workspace: str
    item: str
    item_name: str
    item_type: str
    keep: str
    prior_status: str
    dependency_summary: str
    dependency_evidence: str
    dependency_known: bool
    eligible_for_deletion: bool
    action: str
    detail: str


def _dependency_evidence(lineage_row: tuple[Any, ...] | None, lineage_headers: dict[str, int]) -> tuple[str, str, bool]:
    if lineage_row is None:
        return "", "No matching lineage row", True
    summary = _text(_cell(lineage_row, lineage_headers, "Dependency summary"))
    evidence: list[str] = []
    for name in DEPENDENCY_COLUMNS:
        value = _text(_cell(lineage_row, lineage_headers, name))
        if value:
            evidence.append(f"{name}: {value}")
    total = _cell(lineage_row, lineage_headers, "Impact analysis: total downstream items")
    try:
        if total is not None and int(total) > 0 and not any("Impact" in item for item in evidence):
            evidence.append(f"Impact analysis: total downstream items: {total}")
    except (TypeError, ValueError):
        pass
    known = bool(evidence) or not NO_DEPENDENCY_TEXT.fullmatch(summary or "")
    return summary, " | ".join(evidence), known


def analyze_workbook(path: str | Path) -> tuple[openpyxl.Workbook, list[AuditDecision], dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, data_only=False)
    action_sheet = None
    lineage_sheet = None
    action_headers: dict[str, int] = {}
    lineage_headers: dict[str, int] = {}
    for sheet in workbook.worksheets:
        headers = _header_map(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        # Prefer the action sheet's explicit YES/NO header over the lineage
        # sheet's separate "Recommended to Keep" review field.
        if any("keep?" in name or "yes or no" in name for name in headers):
            action_sheet, action_headers = sheet, headers
        if "dependency summary" in headers:
            lineage_sheet, lineage_headers = sheet, headers
    if action_sheet is None:
        raise ValueError("No worksheet with a Keep header was found")
    if lineage_sheet is None:
        raise ValueError("No worksheet with a Dependency summary header was found")

    keep_key = next(name for name in action_headers if "keep" in name)
    lineage_map = {
        (_text(row[0]), _text(row[1])): row
        for row in lineage_sheet.iter_rows(min_row=2, values_only=True)
        if len(row) >= 2 and _text(row[0]) and _text(row[1])
    }
    decisions: list[AuditDecision] = []
    for row in action_sheet.iter_rows(min_row=2, values_only=True):
        workspace, item = _text(row[0]), _text(row[1]) if len(row) > 1 else ""
        keep = _text(_cell(row, action_headers, keep_key))
        if not keep or not workspace or not item:
            continue
        if keep.casefold() != "no":
            continue
        item_name, item_type = _item_parts(item)
        prior_status = _text(row[5]) if len(row) > 5 else ""
        summary, evidence, known = _dependency_evidence(lineage_map.get((workspace, item)), lineage_headers)
        if prior_status.casefold() == "deleted":
            action, eligible, detail = "ALREADY_MARKED_DELETED", False, "Workbook already marks this item as DELETED"
        elif known:
            action, eligible, detail = "SKIP_DEPENDENCY", False, "Known dependency evidence; deletion skipped"
        else:
            action, eligible, detail = "ELIGIBLE_AFTER_LIVE_RESOLUTION", True, "No known workbook dependency; live identity resolution required"
        decisions.append(AuditDecision(workspace, item, item_name, item_type, keep, prior_status, summary, evidence, known, eligible, action, detail))
    metadata = {"action_sheet": action_sheet.title, "lineage_sheet": lineage_sheet.title, "keep_header": keep_key}
    return workbook, decisions, metadata


def annotate_column_g(workbook: openpyxl.Workbook, decisions: list[AuditDecision], metadata: dict[str, Any]) -> None:
    sheet = workbook[metadata["action_sheet"]]
    sheet.cell(row=1, column=7).value = "Dependency evidence / action"
    by_key = {(decision.workspace, decision.item): decision for decision in decisions}
    for row_number in range(2, sheet.max_row + 1):
        key = (_text(sheet.cell(row_number, 1).value), _text(sheet.cell(row_number, 2).value))
        decision = by_key.get(key)
        if decision is None:
            continue
        if decision.dependency_known:
            text = f"{decision.dependency_summary or 'Known dependency'}; {decision.dependency_evidence}; ACTION={decision.action}"
        else:
            text = f"{decision.dependency_summary or 'No dependencies found'}; ACTION={decision.action}"
        sheet.cell(row=row_number, column=7).value = text[:32767]


class FabricCatalog:
    """Read-only catalog plus explicitly guarded item deletion."""

    API_ROOT = "https://api.fabric.microsoft.com/v1"
    SCOPE = "https://api.fabric.microsoft.com/.default"

    def __init__(self, workspace_aliases: dict[str, str] | None = None) -> None:
        self.workspace_aliases = workspace_aliases or {}
        self._session = None

    def _headers(self) -> dict[str, str]:
        from azure.identity import AzureCliCredential, DefaultAzureCredential, EnvironmentCredential, ManagedIdentityCredential
        if self._session is None:
            token = None
            for credential_type in (AzureCliCredential, ManagedIdentityCredential, EnvironmentCredential, DefaultAzureCredential):
                try:
                    token = credential_type().get_token(self.SCOPE).token
                    break
                except Exception:
                    continue
            if not token:
                raise RuntimeError("No Azure credential could obtain a Fabric API token")
            import requests
            self._session = requests.Session()
            self._session.headers.update({"Authorization": f"Bearer {token}"})
        return dict(self._session.headers)

    def workspaces(self) -> list[dict[str, Any]]:
        import requests
        response = requests.get(f"{self.API_ROOT}/workspaces", headers=self._headers(), timeout=60)
        response.raise_for_status()
        return response.json().get("value", [])

    def resolve_workspace(self, name: str) -> dict[str, Any] | None:
        wanted = self.workspace_aliases.get(name, name)
        return next((item for item in self.workspaces() if item.get("displayName") == wanted or item.get("id") == wanted), None)

    def items(self, workspace_id: str, item_type: str) -> list[dict[str, Any]]:
        import requests
        endpoint = "lakehouses" if item_type.casefold() == "lakehouse" else "warehouses"
        response = requests.get(f"{self.API_ROOT}/workspaces/{workspace_id}/{endpoint}", headers=self._headers(), timeout=60)
        response.raise_for_status()
        return response.json().get("value", [])

    def resolve_item(self, workspace: dict[str, Any], item_name: str, item_type: str) -> dict[str, Any] | None:
        return next((item for item in self.items(workspace["id"], item_type) if item.get("displayName") == item_name), None)

    def delete_item(self, workspace: dict[str, Any], item: dict[str, Any], *, execute: bool) -> str:
        if not execute:
            return "DRY_RUN_DELETE"
        if os.getenv("FABRIC_ALLOW_DELETIONS") != "1" or os.getenv("FABRIC_DELETE_CONFIRMATION") != "DELETE_APPROVED_NO_DEPENDENCY_ITEMS":
            raise RuntimeError("Deletion guard not satisfied; set FABRIC_ALLOW_DELETIONS=1 and the exact confirmation value")
        import requests
        response = requests.delete(f"{self.API_ROOT}/workspaces/{workspace['id']}/items/{item['id']}", headers=self._headers(), timeout=60)
        response.raise_for_status()
        return "DELETED"


def execute_decisions(decisions: list[AuditDecision], *, execute: bool = False, workspace_aliases: dict[str, str] | None = None) -> list[dict[str, Any]]:
    catalog = FabricCatalog(workspace_aliases)
    results: list[dict[str, Any]] = []
    for decision in decisions:
        result = asdict(decision)
        if not decision.eligible_for_deletion:
            result["live_action"] = decision.action
            results.append(result)
            continue
        try:
            workspace = catalog.resolve_workspace(decision.workspace)
            if workspace is None:
                result["live_action"] = "BLOCKED_NO_LIVE_WORKSPACE"
                result["live_detail"] = "Exact workbook workspace was not found; no replacement workspace was guessed"
            else:
                item = catalog.resolve_item(workspace, decision.item_name, decision.item_type)
                if item is None:
                    result["live_action"] = "BLOCKED_NO_LIVE_ITEM"
                    result["live_detail"] = "Exact item name/type was not found in the resolved workspace"
                else:
                    result["live_action"] = catalog.delete_item(workspace, item, execute=execute)
                    result["live_workspace_id"] = workspace.get("id")
                    result["live_item_id"] = item.get("id")
        except Exception as exc:
            result["live_action"] = "ERROR"
            result["live_detail"] = f"{type(exc).__name__}: {exc}"
        results.append(result)
    return results


def run(input_path: str | Path, annotated_path: str | Path, results_path: str | Path, *, execute: bool = False, workspace_aliases: dict[str, str] | None = None) -> list[dict[str, Any]]:
    workbook, decisions, metadata = analyze_workbook(input_path)
    annotate_column_g(workbook, decisions, metadata)
    workbook.save(annotated_path)
    results = execute_decisions(decisions, execute=execute, workspace_aliases=workspace_aliases)
    Path(results_path).write_text(json.dumps(results, indent=2, ensure_ascii=True), encoding="utf-8")
    return results


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--results", default="data/cleanup_results.json")
    parser.add_argument("--execute-deletions", action="store_true")
    parser.add_argument("--workspace-aliases", default="{}", help="JSON object mapping workbook workspace names to exact live display names or IDs")
    args = parser.parse_args()
    aliases = json.loads(args.workspace_aliases)
    results = run(args.input, args.output, args.results, execute=args.execute_deletions, workspace_aliases=aliases)
    counts: dict[str, int] = {}
    for result in results:
        key = str(result.get("live_action", result.get("action")))
        counts[key] = counts.get(key, 0) + 1
    print(json.dumps({"rows": len(results), "actions": counts}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
